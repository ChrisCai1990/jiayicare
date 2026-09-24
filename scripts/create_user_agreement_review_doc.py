from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / 'deliverables' / '嘉医汇健康管家用户协议审核稿.docx'


def font(run, size=10.5, bold=False, color=None):
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement('w:shd')
    node.set(qn('w:fill'), fill)
    props.append(node)


def cell_style(cell, fill=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    props = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for side in ('top', 'start', 'bottom', 'end'):
        item = OxmlElement(f'w:{side}')
        item.set(qn('w:w'), '120')
        item.set(qn('w:type'), 'dxa')
        margins.append(item)
    props.append(margins)
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        item = OxmlElement(f'w:{edge}')
        item.set(qn('w:val'), 'single')
        item.set(qn('w:sz'), '6')
        item.set(qn('w:color'), 'D9E2DE')
        borders.append(item)
    props.append(borders)
    if fill:
        shading(cell, fill)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(15 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    font(r, 14 if level == 1 else 11.5, True)


def paragraph(doc, text, after=7):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.5
    font(p.add_run(text))


def roles_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    headers = ['角色', '主体', '职责边界']
    for i, value in enumerate(headers):
        c = table.rows[0].cells[i]
        cell_style(c, '176B58')
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(value), 9.5, True, (255, 255, 255))
    rows = [
        ('平台运营方', '杭州嘉静佑辰科技有限公司', '运营嘉医汇健康管家平台，提供账号、系统、客服和平台服务。'),
        ('具体服务提供方', '杭州嘉医汇健康管理有限公司', '仅在用户选择相应服务时，提供具体健康管理服务并按页面说明承接付款或开票安排。'),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell_style(cells[i])
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.25
            font(p.add_run(value), 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(2.3), Cm(2.3)
    doc.styles['Normal'].font.name = 'Microsoft YaHei'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    font(p.add_run('嘉医汇健康管家用户协议审核稿'), 22, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    font(p.add_run('版本 2026年9月24日  适用于 App 小程序 官网及相关服务入口'), 10, False, (90, 111, 107))

    paragraph(doc, '本稿用于内部审核。请在实际发布前确认商品页面、支付页面、发票抬头、退款规则及服务单页与本协议的主体和表述保持一致。', 10)
    heading(doc, '一 协议主体与适用范围')
    roles_table(doc)
    paragraph(doc, '本协议适用于用户访问、注册、使用嘉医汇健康管家平台及通过平台了解、预约或购买相关健康管理服务的行为。用户使用平台前应认真阅读本协议；完成注册、登录或继续使用平台，即表示已阅读并同意本协议。')

    heading(doc, '二 平台服务内容')
    paragraph(doc, '平台提供健康档案整理、体检信息整理、健康数据趋势展示、生活方式管理、健康提醒、服务需求梳理、服务预约及就医协助等非医疗健康管理功能。平台不提供疾病诊断、治疗、处方、线上复诊或检查开单。')
    paragraph(doc, '具体服务的内容、适用范围、价格、服务期限、提供方式、退款规则及其他条件，以相应服务页面、订单页面或双方另行确认的服务说明为准。')

    heading(doc, '三 账号注册与使用')
    paragraph(doc, '用户应提供真实、准确、完整的注册和服务信息，并妥善保管账号、验证码和登录凭证。用户不得出租、出借、转让或以其他方式允许他人使用账号。发现账号被盗用或存在安全风险时，应及时通过平台客服反馈。')

    heading(doc, '四 服务提供与授权')
    paragraph(doc, '当用户选择由杭州嘉医汇健康管理有限公司提供具体健康管理服务时，平台将在相应服务页面说明服务提供方、服务内容、费用安排及需要提供的信息。用户确认选择后，相关服务由该服务提供方按约定执行。')
    paragraph(doc, '如为实现用户选择的具体服务确需向服务提供方、健康管理人员或其他经用户授权的接收方提供必要信息，平台将依照隐私政策和服务页面说明处理，并在法律要求时取得用户的单独同意。用户可在适用功能中撤回非必要授权，但撤回可能影响相应服务的继续提供。')

    heading(doc, '五 付款 订单与退款')
    paragraph(doc, '用户提交订单前，应核对服务名称、规格、价格、优惠、支付方式、服务提供方及订单说明。支付成功不当然表示服务已经完成；服务安排、进度和履约记录以平台订单和工作人员确认信息为准。')
    paragraph(doc, '退款、取消及改期应按照服务页面、订单说明和适用法律法规处理。已发生的、经用户确认的第三方费用或已完成的服务部分，按相应规则处理。')

    heading(doc, '六 用户行为规范')
    paragraph(doc, '用户不得利用平台从事违法活动，不得上传虚假或侵害他人权益的信息，不得干扰平台正常运行，不得擅自复制、修改、传播平台内容或程序。')

    heading(doc, '七 隐私与信息保护')
    paragraph(doc, '平台按照《嘉医汇健康管家隐私政策》处理个人信息和健康相关信息。健康相关信息属于敏感个人信息；用户应避免在不支持提交医疗资料的入口提交病历、症状、检查报告、检查指标或用药信息。')

    heading(doc, '八 非医疗服务声明')
    paragraph(doc, '平台及嘉医汇健康管理服务均不替代医疗机构、执业医师或其他依法具备资质的专业人员的诊断、治疗和处方建议。出现身体不适、急症或紧急情况时，请及时就医或拨打120。')

    heading(doc, '九 协议变更与终止')
    paragraph(doc, '平台可在法律法规允许的范围内调整服务。涉及用户权益的重大变更，将通过平台公告、站内消息、短信或其他合理方式提示。用户可依法申请注销账号；法律法规要求保存的信息除外。')

    heading(doc, '十 适用法律与争议处理')
    paragraph(doc, '本协议适用中华人民共和国法律。因本协议产生的争议，双方应先协商解决；协商不成的，可向有管辖权的人民法院提起诉讼。')

    heading(doc, '十一 联系方式')
    paragraph(doc, '平台客服及隐私咨询电话：19106761448。健康管理服务联系地址：杭州市萧山区盈丰街道江峰商务名座1幢1015室。')

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run('嘉医汇健康管家用户协议审核稿'), 8.5, False, (110, 125, 121))
    doc.core_properties.title = '嘉医汇健康管家用户协议审核稿'
    doc.core_properties.author = '杭州嘉静佑辰科技有限公司'
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
