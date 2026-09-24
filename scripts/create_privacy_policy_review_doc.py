from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / 'deliverables' / '嘉医汇健康管家隐私政策审核稿.docx'


def set_font(run, size=None, bold=None, color=None):
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for name, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{name}'))
        if node is None:
            node = OxmlElement(f'w:{name}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_border(cell, color='D9E2DE'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn(f'w:{edge}')
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '6')
        element.set(qn('w:color'), color)


def add_text(doc, text, size=10.5, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    set_font(r, size=size)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    set_font(r, size=14 if level == 1 else 11.5, bold=True, color=(0, 0, 0))
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = 'Table Grid'
    for idx, title in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Cm(widths[idx])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, '176B58')
        set_cell_margins(cell)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        set_font(r, size=9.5, bold=True, color=(255, 255, 255))
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].width = Cm(widths[idx])
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[idx])
            set_cell_border(cells[idx])
            p = cells[idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.25
            r = p.add_run(value)
            set_font(r, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    normal = doc.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run('嘉医汇健康管家隐私政策审核稿')
    set_font(title_run, size=22, bold=True, color=(0, 0, 0))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    r = subtitle.add_run('版本 2026年9月24日  适用于 App 小程序 官网及 AI 咨询入口')
    set_font(r, size=10, color=(90, 111, 107))

    add_text(doc, '本稿用于内部合规审核。平台由杭州嘉静佑辰科技有限公司运营；部分健康管理服务由杭州嘉医汇健康管理有限公司提供。实际发布前，请核实两家主体的对外联系信息、支付与开票安排、第三方服务清单及服务页面的单独同意机制。', after=10)

    add_heading(doc, '一 角色与适用范围')
    add_table(doc,
              ['事项', '主体或范围', '说明'],
              [
                  ('平台运营主体及个人信息处理者', '杭州嘉静佑辰科技有限公司', '负责嘉医汇健康管家平台的账号、系统、客服与个人信息处理。'),
                  ('具体健康管理服务提供方', '杭州嘉医汇健康管理有限公司', '仅在用户选择相应服务时提供具体健康管理服务、付款或开票安排。'),
                  ('适用范围', 'App 小程序 官网 AI 咨询入口', '适用于上述入口中的个人信息处理活动。'),
              ], [4.2, 5.4, 7.0])

    add_text(doc, '嘉医汇健康管家平台重视并保护您的个人信息和健康相关信息。本政策说明我们如何收集、使用、保存和保护您的信息，以及您如何行使相关权利。')

    add_heading(doc, '二 我们收集的信息')
    add_text(doc, '为提供健康档案整理、体检信息整理、健康数据趋势展示、生活方式管理、健康提醒、服务需求梳理及就医协助等服务，我们可能收集您主动提供的注册与基础信息、健康相关信息，以及保障服务运行与安全所需的设备和日志信息。')
    add_text(doc, '健康相关信息属于敏感个人信息。我们仅在实现相应服务所必需的范围内处理，并在需要时取得您的单独同意。')
    add_text(doc, '当您使用官网 AI 咨询准备助手并主动提交线下沟通申请时，我们会在取得您同意后处理姓名、手机号、所在城市、方便联系时间和您填写的非医疗咨询方向或摘要，用于确认健康管理服务安排。该入口不用于接收病历、症状、检查指标、报告或用药信息；请勿在其中提交此类信息。')

    add_heading(doc, '三 信息的使用与服务衔接')
    add_text(doc, '我们在必要范围内使用信息，用于提供和改进健康管理服务、展示健康数据趋势、发送服务通知与提醒、保障账号与服务安全，以及在您明确授权后向相应医生或健康管理人员提供必要信息。我们不会出售您的个人信息或健康相关信息。')
    add_text(doc, '如您选择由杭州嘉医汇健康管理有限公司提供具体健康管理服务，我们将在服务页面告知接收方、处理目的、处理方式和信息类型；在法律要求时取得您的单独同意。您可以按照适用功能的说明撤回授权；撤回不影响撤回前基于授权进行的处理。')

    add_heading(doc, '四 健康信息与专业服务边界')
    add_text(doc, '嘉医汇提供非医疗健康管理服务，不提供疾病诊断、治疗、处方、线上复诊或检查开单。涉及医疗问题时，请以正规医疗机构和执业医师的专业判断为准。')
    add_text(doc, '如您选择向医生或健康管理人员授权共享信息，我们将在授权范围内处理。您可以按照适用功能的说明撤回授权。')

    add_heading(doc, '五 信息存储与安全')
    add_text(doc, '我们采取合理的管理和技术措施保护信息安全，并将按照法律法规和业务所需的期限保存信息。官网线下沟通申请在未转化为正式服务关系的情况下，最长保存180日；超过保存期限或您提出有效删除请求后，我们将按适用法律法规处理相关信息。')

    add_heading(doc, '六 第三方服务')
    add_text(doc, '为实现部分功能，我们可能接入短信验证、云服务、对象存储及依法提供服务的 AI 模型接口。我们会仅在实现对应功能所需的范围内向其提供必要信息，并要求其采取相应的数据保护措施。实际启用的第三方名称、处理目的和信息类型，应在发布前的第三方服务清单中列明。')

    add_heading(doc, '七 您的权利')
    add_text(doc, '在符合法律法规的前提下，您可以请求访问、更正、删除个人信息，撤回同意，申请注销账号，或咨询个人信息处理情况。账号注销申请可在 App 或小程序的帮助与反馈中提交。')

    add_heading(doc, '八 未成年人保护')
    add_text(doc, '嘉医汇目前不面向18周岁以下未成年人提供独立服务。如涉及未成年人信息处理，将依照适用法律法规和监护人授权要求执行。')

    add_heading(doc, '九 政策更新')
    add_text(doc, '如本政策发生重大变化，我们将通过官网、App 或其他适当方式提示您。更新后的政策以公开页面载明的版本和生效日期为准。')

    add_heading(doc, '十 联系我们')
    add_table(doc,
              ['联系事项', '联系信息'],
              [
                  ('平台运营主体及个人信息处理者', '杭州嘉静佑辰科技有限公司'),
                  ('健康管理服务提供方', '杭州嘉医汇健康管理有限公司'),
                  ('平台隐私咨询及客服电话', '19106761448'),
                  ('健康管理服务联系地址', '杭州市萧山区盈丰街道江峰商务名座1幢1015室'),
              ], [6.2, 10.4])

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('嘉医汇健康管家隐私政策审核稿')
    set_font(r, size=8.5, color=(110, 125, 121))

    doc.core_properties.title = '嘉医汇健康管家隐私政策审核稿'
    doc.core_properties.author = '杭州嘉静佑辰科技有限公司'
    doc.core_properties.subject = '隐私政策内部审核稿'
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
